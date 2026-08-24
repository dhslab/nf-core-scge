#!/usr/bin/env python3
"""Compile the per-sample SCGE report JSON into a multi-sheet Excel workbook.

Engine note: this runs on whichever of openpyxl / xlsxwriter the container provides.
docker-baseimage ships pandas + openpyxl + Pillow and NOT xlsxwriter, so openpyxl is the
path that actually executes; xlsxwriter is still supported if present. The two libraries
have different worksheet APIs (create_sheet/cell/add_image vs add_worksheet/write/
insert_image), so every call below goes through a small adapter rather than being written
against one of them with the other bolted on.
"""

import argparse
import json
import os
import sys

import pandas as pd


def parse_args():
    parser = argparse.ArgumentParser(description='Generate Excel report for SCGE pipeline')
    parser.add_argument('--report_json', required=True, help='Path to report JSON file')
    parser.add_argument('--circos_plot', help='Path to Circos plot image')
    parser.add_argument('--cna_plot', help='Path to CNA plot image')
    parser.add_argument('--baf_plot', help='Path to BAF plot image')
    parser.add_argument('--indel_freq_plot', help='Path to Indel Frequency plot image')
    parser.add_argument('--off_targets_plot', help='Path to Off-targets plot image')
    parser.add_argument('--output', required=True, help='Output Excel filename')
    return parser.parse_args()


def usable(path):
    """A plot argument is usable only if it was passed, exists and is non-empty.

    Upstream processes `touch` a placeholder when they have nothing to draw, so an empty
    file is the normal 'no plot for this sample' signal, not an error.
    """
    return bool(path) and os.path.exists(path) and os.path.getsize(path) > 0


class SummarySheet:
    """Minimal write/insert_image surface over either Excel engine."""

    def __init__(self, workbook, engine, title='Summary'):
        self.engine = engine
        if engine == 'xlsxwriter':
            self.ws = workbook.add_worksheet(title)
            self.bold = workbook.add_format({'bold': True})
            self.title_fmt = workbook.add_format({'bold': True, 'font_size': 14})
        else:
            from openpyxl.styles import Font
            self.ws = workbook.create_sheet(title)
            self.bold = Font(bold=True)
            self.title_fmt = Font(bold=True, size=14)

    def write(self, row, col, value, style=None):
        """Write a cell using 0-based (row, col), as xlsxwriter does throughout."""
        if self.engine == 'xlsxwriter':
            self.ws.write(row, col, value, style)
        else:
            cell = self.ws.cell(row=row + 1, column=col + 1, value=value)
            if style is not None:
                cell.font = style

    def insert_image(self, path, row, col, scale=0.5):
        if not usable(path):
            return False
        try:
            if self.engine == 'xlsxwriter':
                from xlsxwriter.utility import xl_rowcol_to_cell
                self.ws.insert_image(xl_rowcol_to_cell(row, col), path,
                                     {'x_scale': scale, 'y_scale': scale})
            else:
                from openpyxl.drawing.image import Image
                from openpyxl.utils import get_column_letter
                img = Image(path)
                # openpyxl has no scale factor; resize explicitly so a full-width plot does
                # not swamp the sheet the way it would at native resolution.
                img.width = int(img.width * scale)
                img.height = int(img.height * scale)
                self.ws.add_image(img, f"{get_column_letter(col + 1)}{row + 1}")
            return True
        except Exception as e:                                   # noqa: BLE001
            print(f"Warning: could not insert image {path}: {e}")
            return False


def main():
    args = parse_args()

    try:
        with open(args.report_json, 'r') as f:
            data = json.load(f)
    except Exception as e:                                       # noqa: BLE001
        print(f"Error loading JSON: {e}")
        sys.exit(1)

    try:
        import openpyxl                                          # noqa: F401
        engine = 'openpyxl'
    except ImportError:
        try:
            import xlsxwriter                                    # noqa: F401
            engine = 'xlsxwriter'
        except ImportError:
            print("Error: neither openpyxl nor xlsxwriter is installed.")
            sys.exit(1)

    writer = pd.ExcelWriter(args.output, engine=engine)
    workbook = writer.book

    # --- Sheet 1: summary and plots ---
    summary = SummarySheet(workbook, engine)
    summary.write(0, 0, 'Somatic Editing Genome Report', summary.title_fmt)

    meta = data.get('metadata', {})
    coverage = meta.get('mean_coverage', {})
    metadata = [
        ('Sample ID', data.get('sample_id', 'N/A')),
        ('Control Sample', meta.get('control_sample', 'N/A')),
        ('Transgene', data.get('transgene_description', 'N/A')),
        ('Tumor Mean Coverage', f"{coverage.get('tumor', 'N/A')}x"),
        ('Normal Mean Coverage', f"{coverage.get('normal', 'N/A')}x"),
    ]
    row = 2
    for key, val in metadata:
        summary.write(row, 0, key, summary.bold)
        summary.write(row, 1, val)
        row += 1

    # Two columns of plots (A and I), each block given room for a half-scale image.
    row += 2
    for label, path, col in [
        ('On-Target Indel Frequency', args.indel_freq_plot, 0),
        ('Transgene Integrations (Circos)', args.circos_plot, 8),
    ]:
        summary.write(row, col, label, summary.bold)
        summary.insert_image(path, row + 2, col)

    row += 25
    summary.write(row, 0, 'Off-Target Sites', summary.bold)
    summary.insert_image(args.off_targets_plot, row + 2, 0)

    row += 25
    for label, path, col in [
        ('Copy Number Analysis (CNA)', args.cna_plot, 0),
        ('B-Allele Frequency (BAF)', args.baf_plot, 8),
    ]:
        summary.write(row, col, label, summary.bold)
        summary.insert_image(path, row + 2, col)

    # --- Sheets 2-4: the tables ---
    tables = data.get('tables', {})

    on_target = tables.get('on_target_sv_transgene', [])
    if on_target:
        df = pd.DataFrame(on_target if isinstance(on_target, list) else [on_target])
        cols = [c for c in ['Location', 'SYMBOL', 'Consequence', 'EXON', 'INTRON',
                            'HGVSc', 'HGVSp'] if c in df.columns]
        if cols:
            df = df[cols]
    else:
        df = pd.DataFrame({'Message': ['No on-target variants found']})
    df.to_excel(writer, sheet_name='On-Target Variants', index=False)

    off_target = tables.get('off_target_indels', [])
    (pd.DataFrame(off_target) if off_target
     else pd.DataFrame({'Message': ['No off-target indels found']})
     ).to_excel(writer, sheet_name='Off-Target Indels', index=False)

    targeted = tables.get('targeted_gene_mutations', [])
    (pd.DataFrame(targeted) if targeted
     else pd.DataFrame({'Message': ['No targeted gene mutation data available']})
     ).to_excel(writer, sheet_name='Targeted Gene Mutations', index=False)

    writer.close()
    print(f"Successfully generated {args.output} (engine: {engine})")


if __name__ == "__main__":
    main()
